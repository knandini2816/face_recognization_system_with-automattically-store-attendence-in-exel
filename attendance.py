import sqlite3
import os
from datetime import datetime, time
from openpyxl import Workbook, load_workbook

DATABASE = "database/attendance.db"

REPORT_FOLDER = "reports"
EXCEL_FILE = os.path.join(
    REPORT_FOLDER,
    "attendance.xlsx"
)

os.makedirs(REPORT_FOLDER, exist_ok=True)


# --------------------------------------------------
# CREATE DATABASE TABLES
# --------------------------------------------------

def create_table():

    connection = sqlite3.connect(DATABASE)
    cursor = connection.cursor()

    cursor.execute("""
        CREATE TABLE IF NOT EXISTS attendance (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            student_id TEXT,
            roll_number TEXT,
            name TEXT,
            department TEXT,
            year TEXT,
            date TEXT,
            time TEXT,
            session TEXT,
            capture_count INTEGER DEFAULT 1
        )
    """)

    cursor.execute("""
        CREATE TABLE IF NOT EXISTS capture_attempts (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            student_id TEXT,
            date TEXT,
            session TEXT
        )
    """)

    connection.commit()
    connection.close()


# --------------------------------------------------
# GET CURRENT SESSION
# --------------------------------------------------

def get_current_session():

    current_time = datetime.now().time()

    if time(9, 0) <= current_time <= time(12, 0):
        return "Morning"

    if time(13, 0) <= current_time <= time(16, 0):
        return "Afternoon"

    return None


# --------------------------------------------------
# GET CAPTURE COUNT
# --------------------------------------------------

def get_capture_count(student_id):

    today = datetime.now().strftime("%Y-%m-%d")
    session = get_current_session()

    if session is None:
        return 0

    connection = sqlite3.connect(DATABASE)
    cursor = connection.cursor()

    cursor.execute("""
        SELECT COUNT(*)
        FROM capture_attempts
        WHERE student_id = ?
        AND date = ?
        AND session = ?
    """, (
        student_id,
        today,
        session
    ))

    result = cursor.fetchone()

    connection.close()

    return result[0]


# --------------------------------------------------
# UPDATE CAPTURE COUNT IN EXCEL
# --------------------------------------------------

def update_capture_count_in_excel(
    student_id,
    date,
    session,
    capture_count
):

    # Create Excel if it does not exist
    if not os.path.exists(EXCEL_FILE):

        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = "Attendance"

        worksheet.append([
            "Student ID",
            "Roll Number",
            "Name",
            "Department",
            "Year",
            "Date",
            "Time",
            "Session",
            "Capture Count"
        ])

        workbook.save(EXCEL_FILE)
        workbook.close()

        return

    workbook = load_workbook(EXCEL_FILE)
    worksheet = workbook["Attendance"]

    headers = [
        cell.value
        for cell in worksheet[1]
    ]

    # Find Capture Count column
    if "Capture Count" not in headers:

        capture_column = len(headers) + 1

        worksheet.cell(
            row=1,
            column=capture_column,
            value="Capture Count"
        )

    else:

        capture_column = headers.index(
            "Capture Count"
        ) + 1

    # Find existing student/session row
    found_row = None

    for row in range(
        2,
        worksheet.max_row + 1
    ):

        existing_student_id = worksheet.cell(
            row=row,
            column=1
        ).value

        existing_date = worksheet.cell(
            row=row,
            column=6
        ).value

        existing_session = worksheet.cell(
            row=row,
            column=8
        ).value

        if (
            str(existing_student_id) == str(student_id)
            and str(existing_date) == str(date)
            and str(existing_session) == str(session)
        ):

            found_row = row
            break

    # Update existing row
    if found_row is not None:

        worksheet.cell(
            row=found_row,
            column=capture_column,
            value=capture_count
        )

    workbook.save(EXCEL_FILE)
    workbook.close()


# --------------------------------------------------
# RECORD A NEW CAPTURE
# --------------------------------------------------

def record_capture(student_id):

    today = datetime.now().strftime("%Y-%m-%d")
    session = get_current_session()

    if session is None:
        return 0

    connection = sqlite3.connect(DATABASE)
    cursor = connection.cursor()

    # Add new capture
    cursor.execute("""
        INSERT INTO capture_attempts (
            student_id,
            date,
            session
        )
        VALUES (?, ?, ?)
    """, (
        student_id,
        today,
        session
    ))

    connection.commit()

    # Get updated count
    cursor.execute("""
        SELECT COUNT(*)
        FROM capture_attempts
        WHERE student_id = ?
        AND date = ?
        AND session = ?
    """, (
        student_id,
        today,
        session
    ))

    result = cursor.fetchone()

    connection.close()

    capture_count = result[0]

    # IMPORTANT:
    # Update Excel immediately
    update_capture_count_in_excel(
        student_id,
        today,
        session,
        capture_count
    )

    return capture_count


# --------------------------------------------------
# CHECK ATTENDANCE
# --------------------------------------------------

def already_marked(student_id):

    today = datetime.now().strftime("%Y-%m-%d")
    session = get_current_session()

    if session is None:
        return False

    connection = sqlite3.connect(DATABASE)
    cursor = connection.cursor()

    cursor.execute("""
        SELECT id
        FROM attendance
        WHERE student_id = ?
        AND date = ?
        AND session = ?
    """, (
        student_id,
        today,
        session
    ))

    result = cursor.fetchone()

    connection.close()

    return result is not None


# --------------------------------------------------
# PREPARE EXCEL
# --------------------------------------------------

def prepare_excel():

    os.makedirs(
        REPORT_FOLDER,
        exist_ok=True
    )

    if not os.path.exists(EXCEL_FILE):

        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = "Attendance"

        worksheet.append([
            "Student ID",
            "Roll Number",
            "Name",
            "Department",
            "Year",
            "Date",
            "Time",
            "Session",
            "Capture Count"
        ])

        workbook.save(EXCEL_FILE)
        workbook.close()

    else:

        workbook = load_workbook(EXCEL_FILE)
        worksheet = workbook["Attendance"]

        headers = [
            cell.value
            for cell in worksheet[1]
        ]

        if "Capture Count" not in headers:

            worksheet.cell(
                row=1,
                column=len(headers) + 1,
                value="Capture Count"
            )

        workbook.save(EXCEL_FILE)
        workbook.close()


# --------------------------------------------------
# SAVE ATTENDANCE TO EXCEL
# --------------------------------------------------

def save_to_excel(
    student,
    date,
    current_time,
    session,
    capture_count
):

    prepare_excel()

    workbook = load_workbook(EXCEL_FILE)
    worksheet = workbook["Attendance"]

    headers = [
        cell.value
        for cell in worksheet[1]
    ]

    capture_column = headers.index(
        "Capture Count"
    ) + 1

    found_row = None

    for row in range(
        2,
        worksheet.max_row + 1
    ):

        existing_student_id = worksheet.cell(
            row=row,
            column=1
        ).value

        existing_date = worksheet.cell(
            row=row,
            column=6
        ).value

        existing_session = worksheet.cell(
            row=row,
            column=8
        ).value

        if (
            str(existing_student_id) == str(student["student_id"])
            and str(existing_date) == str(date)
            and str(existing_session) == str(session)
        ):

            found_row = row
            break

    # Existing attendance → update count
    if found_row is not None:

        worksheet.cell(
            row=found_row,
            column=capture_column,
            value=capture_count
        )

    else:

        new_row = worksheet.max_row + 1

        values = [
            student["student_id"],
            student["roll_number"],
            student["name"],
            student["department"],
            student["year"],
            date,
            current_time,
            session,
            capture_count
        ]

        for column, value in enumerate(
            values,
            start=1
        ):

            worksheet.cell(
                row=new_row,
                column=column,
                value=value
            )

    workbook.save(EXCEL_FILE)
    workbook.close()


# --------------------------------------------------
# MARK ATTENDANCE
# --------------------------------------------------

def mark_attendance(
    student,
    capture_count
):

    session = get_current_session()

    if session is None:
        return False

    # Do not mark attendance twice
    if already_marked(
        student["student_id"]
    ):
        return False

    now = datetime.now()

    date = now.strftime(
        "%Y-%m-%d"
    )

    current_time = now.strftime(
        "%H:%M:%S"
    )

    connection = sqlite3.connect(
        DATABASE
    )

    cursor = connection.cursor()

    cursor.execute("""
        INSERT INTO attendance (
            student_id,
            roll_number,
            name,
            department,
            year,
            date,
            time,
            session,
            capture_count
        )
        VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
    """, (
        student["student_id"],
        student["roll_number"],
        student["name"],
        student["department"],
        student["year"],
        date,
        current_time,
        session,
        capture_count
    ))

    connection.commit()
    connection.close()

    save_to_excel(
        student,
        date,
        current_time,
        session,
        capture_count
    )

    return True


# --------------------------------------------------
# CREATE TABLES WHEN PROGRAM STARTS
# --------------------------------------------------

create_table()